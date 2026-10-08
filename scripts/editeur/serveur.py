"""
Éditeur local des fiches (data/frise-philosophes-*.xlsx) et des courants de pensée
(data/frise-courant-pensee-*.xlsx), au choix par le sélecteur en haut de la page.
Pour un courant : les philosophes rattachés (ceux dont la colonne Courant le cite) se
rattachent ou se retirent depuis sa page, et un nouveau nom de courant est reporté dans
leurs fiches.

Lancement : double-clic sur « Editeur des fiches.bat » à la racine de laphilo-astro
(ou : python scripts/editeur/serveur.py). La page s'ouvre dans le navigateur.

- L'éditeur écrit dans les xlsx (et les images dans public/pho/) ; rien n'est mis
  en ligne avant le bouton « Publier » (publier.py).
- L'écriture passe par scripts/xlsx_cellules.set_cells : seules les cellules
  modifiées changent, la mise en forme du classeur est conservée.
- Une copie du xlsx est faite avant chaque enregistrement dans
  ../groupes/sauvegarde/editeur/.
- Un xlsx ouvert dans Excel ne peut pas être modifié : l'éditeur le signale.
"""

import base64
import datetime
import glob
import html
import importlib.util
import io
import json
import os
import re
import shutil
import subprocess
import sys
import threading
import urllib.parse
import urllib.request
import webbrowser
import zipfile
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import openpyxl
from openpyxl.utils import get_column_letter
from PIL import Image, ImageOps

ICI = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.dirname(os.path.dirname(ICI))
DATA = os.path.join(BASE, 'data')
PHO = os.path.join(BASE, 'public', 'pho')
PHOTO = os.path.join(BASE, 'public', 'photo')  # grandes photos (colonne Image_media) quand il n'y a pas de vidéo
SAUVEGARDES = os.path.join(os.path.dirname(BASE), 'groupes', 'sauvegarde', 'editeur')
PORT = 8765
# Version du code : un éditeur relancé après une mise à jour remplace l'ancien encore ouvert
def version_fichiers():
    return str(max(os.path.getmtime(os.path.join(ICI, n)) for n in ('serveur.py', 'index.html', 'liens.py', 'publier.py', 'dates.py', 'statsgoogle.py')))


VERSION = version_fichiers()
_relance = threading.Event()


def relancer_si_mis_a_jour():
    """Code de l'éditeur modifié depuis le lancement : un nouvel éditeur démarre (dans sa propre
    fenêtre noire) et remplace celui-ci (il lui demande de s'arrêter). True si la relance est lancée."""
    if version_fichiers() == VERSION:
        return False
    if not _relance.is_set():
        _relance.set()
        args = [sys.executable, os.path.abspath(__file__), '--sans-navigateur'] + (['--port', str(PORT)] if PORT != 8765 else [])
        subprocess.Popen(args, cwd=BASE, creationflags=getattr(subprocess, 'CREATE_NEW_CONSOLE', 0))
    return True


PAGE_RELANCE = '''<!doctype html><html lang="fr"><head><meta charset="utf-8"><title>Mise à jour de l'éditeur…</title>
<meta http-equiv="refresh" content="4"><style>body{margin:0;height:100vh;display:grid;place-items:center;background:#1a0f07;color:#f6ecd0;font-family:Georgia,serif}</style></head>
<body><p>L'éditeur a été mis à jour : redémarrage… (la page se recharge toute seule)</p></body></html>'''

sys.path.insert(0, os.path.join(BASE, 'scripts'))
from xlsx_cellules import set_cells, first_sheet_path, col_num  # noqa: E402
sys.path.insert(0, ICI)
import liens  # noqa: E402  contrôle des liens morts
import publier  # noqa: E402  bouton « Publier »
import dates  # noqa: E402  contrôle des dates par Wikidata (philosophes vivants décédés…)
import statsgoogle  # noqa: E402  bouton « Google » : chiffres Search Console et indexation des fiches

_spec = importlib.util.spec_from_file_location('build_data', os.path.join(BASE, 'scripts', 'build-data.py'))
bd = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(bd)

BRANCHES = ['metaphysique', 'epistemologie', 'ethique', 'politique', 'spiritualite', 'esthetique', 'logique', 'langage']
TRADITIONS = ['juive', 'femme', 'non-binaire'] + list(bd.GROUPES)
PREFIXE_IMAGE = 'https://www.laphilo.fr/pho/'
PREFIXE_PHOTO = 'https://www.laphilo.fr/photo/'
# Images sources (jpg déjà recadrés par l'utilisateur), converties en webp par l'éditeur
SOURCES_IMAGES = os.path.join(os.path.dirname(BASE), 'pho-jpg')
DOMAINES = os.path.join(BASE, 'public', 'data', 'iframe-blocked-domains.json')
verrou = threading.Lock()


# ── Lecture des xlsx (en cache tant que le fichier ne change pas) ──────────
_cache = {}


def chemin(fichier):
    if fichier not in bd.PHILOSOPHE_FICHIERS and fichier not in bd.COURANT_FICHIERS:
        raise ValueError(f'Fichier inconnu : {fichier}')
    return os.path.join(DATA, fichier + '.xlsx')


def est_courant(fichier):
    return fichier in bd.COURANT_FICHIERS


def classeur(fichier):
    p = chemin(fichier)
    mtime = os.path.getmtime(p)
    c = _cache.get(fichier)
    if c and c['mtime'] == mtime:
        return c
    ws = openpyxl.load_workbook(p, data_only=True).worksheets[0]
    entetes = [c.value for c in ws[3]]
    while entetes and entetes[-1] is None:
        entetes.pop()
    lignes = {}
    for r in range(4, ws.max_row + 1):
        valeurs = [ws.cell(r, i + 1).value for i in range(len(entetes))]
        if any(v not in (None, '') for v in valeurs[:5]):
            lignes[r] = valeurs
    c = {'mtime': mtime, 'entetes': entetes, 'lignes': lignes}
    _cache[fichier] = c
    return c


def norm(h):
    return bd.normalize_header(h) if h else ''


def valeur(c, valeurs, cle):
    for i, h in enumerate(c['entetes']):
        if norm(h) == cle:
            return valeurs[i]
    return None


def image_locale(url):
    s = bd.url_relative(url or '')
    for prefixe, dossier in (('/pho/', PHO), ('/photo/', PHOTO)):
        if s.startswith(prefixe):
            return os.path.exists(os.path.join(dossier, s[len(prefixe):])), s
    return bool(s), s


def mesure_texte(h):
    """Longueur du texte lu (sans balises ni bloc audio, espaces réduits) et nombre de liens
    complémentaires (le lien Wikipédia du nom, en tête, n'est pas compté). Même calcul que mesureTexte() dans index.html."""
    h = re.sub(r'<audio.*?</audio>', '', str(h or ''), flags=re.S | re.I)
    liens = re.findall(r'<a\s[^>]*?href="([^"]*)"', h, flags=re.I)
    if liens and 'wikipedia.org' in liens[0]:
        liens = liens[1:]
    texte = re.sub(r'\s+', ' ', html.unescape(re.sub(r'<[^>]+>', ' ', h))).strip()
    return len(texte), len(liens)


def defauts_ligne(c, v):
    """(liens morts, liens à vérifier) d'une ligne, d'après le dernier contrôle des liens."""
    d = liens.defauts_de(valeur(c, v, 'texte_html'), valeur(c, v, 'youtube_id'))
    return sum(1 for x in d if x['statut'] == 'MORT'), sum(1 for x in d if x['statut'] != 'MORT')


def toutes_les_fiches():
    """[(texte, yt_id)] de toutes les fiches, philosophes et courants (pour le contrôle complet des liens)."""
    out = []
    for f in list(bd.PHILOSOPHE_FICHIERS) + list(bd.COURANT_FICHIERS):
        c = classeur(f)
        for v in c['lignes'].values():
            if valeur(c, v, 'nom'):
                out.append((valeur(c, v, 'texte_html'), valeur(c, v, 'youtube_id')))
    return out


def anomalies_ligne(c, v):
    """Écarts de dates avec Wikidata (dernier contrôle « Dates ») d'une ligne de philosophe."""
    nom = valeur(c, v, 'nom')
    return dates.anomalies(bd.slugify(str(nom)), bd.to_int(valeur(c, v, 'annee_naissance')),
                           bd.to_int(valeur(c, v, 'annee_deces')), valeur(c, v, 'dates_affichage')) if nom else []


def philosophes_slug_texte():
    out = []
    for f in bd.PHILOSOPHE_FICHIERS:
        c = classeur(f)
        for v in c['lignes'].values():
            nom = valeur(c, v, 'nom')
            if nom:
                out.append((bd.slugify(str(nom)), valeur(c, v, 'texte_html')))
    return out


def page_fiche(slug, courant=False):
    """Chemin de la page publiée d'une fiche (clé des statistiques Google)."""
    return f"/{'courants' if courant else 'philosophes'}/{slug}/"


def google_ligne(slug, courant=False):
    """Pour la liste : affichages des 90 derniers jours et état d'indexation."""
    g = statsgoogle.etat()
    ch = page_fiche(slug, courant)
    x = g['pages'].get(ch) or {}
    return {'g_i': x.get('i', 0), 'g_c': x.get('c', 0), 'g_idx': statsgoogle.categorie(g['index'].get(ch))}


def toutes_les_pages():
    """[chemin] de toutes les fiches, philosophes puis courants (pour le contrôle d'indexation)."""
    out = []
    for f in list(bd.PHILOSOPHE_FICHIERS) + list(bd.COURANT_FICHIERS):
        c = classeur(f)
        for v in c['lignes'].values():
            nom = valeur(c, v, 'nom')
            if nom:
                out.append(page_fiche(bd.slugify(str(nom)), est_courant(f)))
    return list(dict.fromkeys(out))


def liste_google():
    """Bilan du bouton « Google » : indexation par état, fiches les plus vues, autres pages du site."""
    g = statsgoogle.etat()
    fiches, chemins = [], set()
    for f in list(bd.PHILOSOPHE_FICHIERS) + list(bd.COURANT_FICHIERS):
        c = classeur(f)
        for r, v in c['lignes'].items():
            nom = valeur(c, v, 'nom')
            if not nom:
                continue
            ch = page_fiche(bd.slugify(str(nom)), est_courant(f))
            chemins.add(ch)
            x = g['pages'].get(ch) or {}
            fiches.append({'f': f, 'l': r, 'nom': str(nom), 'courant': est_courant(f), 'chemin': ch,
                           'i': x.get('i', 0), 'c': x.get('c', 0), 'p': x.get('p'),
                           'idx': statsgoogle.categorie(g['index'].get(ch))})
    par_etat = {}
    for x in fiches:
        par_etat[x['idx']] = par_etat.get(x['idx'], 0) + 1
    vues = sorted([x for x in fiches if x['i']], key=lambda x: (-x['c'], -x['i']))
    autres = sorted(([ch, x['i'], x['c'], x['p']] for ch, x in g['pages'].items() if ch not in chemins),
                    key=lambda y: (-y[2], -y[1]))
    return {**statsgoogle.resume(), 'total': len(fiches), 'par_etat': par_etat, 'libelles': statsgoogle.LIBELLES,
            'vues': vues, 'autres': autres[:40],
            'totaux': {'i': sum(x['i'] for x in fiches), 'c': sum(x['c'] for x in fiches)}}


def liste_dates():
    """Fiches de philosophes dont les dates diffèrent de Wikidata, décédés d'abord."""
    out, sans = [], []
    etat_sans = dates.etat().get('sans', {})
    for f in bd.PHILOSOPHE_FICHIERS:
        c = classeur(f)
        for r, v in c['lignes'].items():
            nom = valeur(c, v, 'nom')
            if not nom:
                continue
            a = anomalies_ligne(c, v)
            if a:
                out.append({'f': f, 'l': r, 'nom': str(nom), 'dates': str(valeur(c, v, 'dates_affichage') or ''), 'anomalies': a})
            raison = etat_sans.get(bd.slugify(str(nom)))
            if raison:
                sans.append({'f': f, 'l': r, 'nom': str(nom), 'raison': raison})
    out.sort(key=lambda x: (not any(a['type'] == 'decede' for a in x['anomalies']), bd.normalize_header(x['nom'])))
    return {**dates.resume(), 'fiches': out, 'non_comparees': sans}


# ── Déplacer une fiche vers un autre fichier (philosophe décédé : fichier « actif » -> historique) ──
def cle_entete(h):
    return bd.normalize_header(h) if h else ''


def retirer_ligne(p, n):
    """Supprime la ligne n et remonte les suivantes (cellules, liens hypertexte) sans réécrire le classeur."""
    tmp = p + '.tmp'
    with zipfile.ZipFile(p) as zi:
        feuille = first_sheet_path(zi)
        xml = zi.read(feuille).decode('utf-8')
        m = re.search(r'<row [^>]*?\br="%d"[^>]*?(/>|>.*?</row>)' % n, xml, re.S)
        if not m:
            raise ValueError(f'ligne {n} introuvable')
        xml = xml[:m.start()] + xml[m.end():]
        xml = re.sub(r'<hyperlink [^>]*?\bref="[A-Z]+%d"[^>]*/>' % n, '', xml)
        xml = re.sub(r'<hyperlinks>\s*</hyperlinks>', '', xml)

        def row(mm):
            r = int(mm.group(2))
            return mm.group(1) + str(r - 1 if r > n else r) + mm.group(3)
        xml = re.sub(r'(<row [^>]*?\br=")(\d+)(")', row, xml)
        xml = re.sub(r'(<c r="[A-Z]+)(\d+)(")', row, xml)
        xml = re.sub(r'(<hyperlink [^>]*?\bref="[A-Z]+)(\d+)(")', row, xml)
        with zipfile.ZipFile(tmp, 'w') as zo:
            for it in zi.infolist():
                zo.writestr(it, xml.encode('utf-8') if it.filename == feuille else zi.read(it.filename))
    shutil.move(tmp, p)


def fichier_historique_propose(f, l):
    """Fichier conseillé pour un philosophe d'un fichier « actif » qui vient de mourir."""
    if f == 'frise-philosophes-france-actif':
        return 'frise-philosophes-france'
    c = classeur(f)
    v = c['lignes'][l]
    nat = bd.normalize_header(valeur(c, v, 'nationalite') or '')
    groupe = valeur(c, v, 'groupe') or ''
    if any(x in nat for x in ('allemand', 'autrichien')):
        return 'frise-philosophes-allemands'
    if any(x in nat for x in ('russe', 'ukrainien', 'polonais', 'tcheque', 'hongrois', 'roumain', 'bulgare', 'serbe')):
        return 'frise-philosophes-russes'
    if any(x in nat for x in ('americain', 'canadien', 'bresil', 'argentin', 'mexicain', 'chilien', 'peruvien',
                              'colombien', 'cubain', 'venezuel', 'uruguay', 'haitien', 'martiniquais')):
        return 'frise-philosophes-americains'
    if groupe and groupe not in ('occident', 'russe'):
        return 'frise-philosophes-orientaux'
    return 'frise-philosophes-modernes'


def deplacer_fiche(f, l, cible, deces=None):
    """Déplace la fiche vers le fichier cible (fin du fichier, mise en forme de la dernière fiche, puis rangée
    par année), en inscrivant l'année de décès si elle est donnée. Renvoie la ligne dans le fichier cible."""
    if f not in bd.PHILOSOPHE_FICHIERS or cible not in bd.PHILOSOPHE_FICHIERS or f == cible:
        raise ValueError('Fichier de destination invalide')
    ps, pc = chemin(f), chemin(cible)
    c = classeur(f)
    if l not in c['lignes']:
        raise ValueError('Ligne introuvable')
    v = c['lignes'][l]
    valeurs = {cle_entete(h): v[i] for i, h in enumerate(c['entetes']) if h}
    n = bd.to_int(valeurs.get('annee_naissance'))
    m = bd.to_int(deces) if str(deces or '').strip() else None
    if m is not None:
        valeurs['annee_deces'] = str(m)
        if n is not None:
            valeurs['dates_affichage'] = dates_affichees(n, m)
    if cible not in bd.FICHIERS_VIVANTS:
        valeurs.pop('themes', None)  # sujets de travail : seulement pour les vivants
    ct = classeur(cible)
    L = lignes_annees(cible)
    pleines = [x[0] for x in L if x[2]]
    vides = [x[0] for x in L if not x[2]]
    modele = pleines[-1] if pleines else 4
    ligne = vides[0] if vides else (max([x[0] for x in L] or [3]) + 1)
    cellules, styles, perdues = {}, {}, []
    index_cible = {cle_entete(h): i for i, h in enumerate(ct['entetes']) if h}
    for k, val in valeurs.items():
        if val in (None, ''):
            continue
        if k not in index_cible:
            perdues.append(k)
            continue
        col = get_column_letter(index_cible[k] + 1)
        cellules[f'{col}{ligne}'] = val
        styles[f'{col}{ligne}'] = f'{col}{modele}'
    with verrou:
        for p in (ps, pc):
            if ouvert_dans_excel(p):
                raise PermissionError(f'{os.path.basename(p)} est ouvert dans Excel : fermez-le puis recommencez.')
        sauvegarder(ps); sauvegarder(pc)
        set_cells(pc, cellules, style_from=styles)
        _cache.pop(cible, None)
        retirer_ligne(ps, l)
        _cache.pop(f, None)
    return placer_ligne(cible, ligne), perdues


def liste_liens_morts():
    """Toutes les fiches qui ont au moins un lien en défaut, philosophes et courants."""
    out = []
    for f in list(bd.PHILOSOPHE_FICHIERS) + list(bd.COURANT_FICHIERS):
        c = classeur(f)
        for r, v in c['lignes'].items():
            nom = valeur(c, v, 'nom')
            d = liens.defauts_de(valeur(c, v, 'texte_html'), valeur(c, v, 'youtube_id')) if nom else []
            if d:
                out.append({'f': f, 'l': r, 'nom': str(nom), 'courant': est_courant(f), 'defauts': d})
    out.sort(key=lambda x: (-sum(1 for d in x['defauts'] if d['statut'] == 'MORT'), bd.normalize_header(x['nom'])))
    return {**liens.resume(), 'fiches': out}


def decouper_courants(s):
    return [x.strip() for x in str(s or '').replace('\xa0', ' ').split(';') if x.strip()]


def rattaches():
    """nom de courant -> [{f, l, nom, dates}] des philosophes qui le citent dans leur colonne Courant."""
    out = {}
    for f in bd.PHILOSOPHE_FICHIERS:
        c = classeur(f)
        for r, v in c['lignes'].items():
            nom = valeur(c, v, 'nom')
            for cour in decouper_courants(valeur(c, v, 'courant')) if nom else []:
                out.setdefault(cour, []).append({'f': f, 'l': r, 'nom': str(nom), 'dates': str(valeur(c, v, 'dates_affichage') or ''),
                                                 'annee': bd.to_int(valeur(c, v, 'annee_naissance'))})
    for L in out.values():
        L.sort(key=lambda x: x['annee'] if x['annee'] is not None else 10 ** 9)
    return out


def liste_courants():
    out = []
    R = rattaches()
    for f in bd.COURANT_FICHIERS:
        c = classeur(f)
        for r, v in c['lignes'].items():
            nom = valeur(c, v, 'nom')
            if not nom:
                continue
            photo_ok, _ = image_locale(valeur(c, v, 'image_media'))
            lg, nb_liens = mesure_texte(valeur(c, v, 'texte_html'))
            morts, douteux = defauts_ligne(c, v)
            out.append({
                'lg': lg, 'liens': nb_liens, 'morts': morts, 'douteux': douteux, 'courant': True,
                'f': f, 'l': r, 'nom': str(nom), 'slug': bd.slugify(str(nom)), 'dates': str(valeur(c, v, 'dates_affichage') or ''),
                'groupe': valeur(c, v, 'groupe') or '', 'image': '', 'image_prevue': '',
                'video': bool(valeur(c, v, 'youtube_id')), 'photo': bool(valeur(c, v, 'image_media')) and photo_ok,
                'annee': bd.to_int(valeur(c, v, 'annee_naissance')), 'nb_phi': len(R.get(str(nom).strip(), [])),
                **google_ligne(bd.slugify(str(nom)), True),
            })
    return out


def liste_fiches(type_='philosophes'):
    if type_ == 'courants':
        return liste_courants()
    out = []
    for f in bd.PHILOSOPHE_FICHIERS:
        c = classeur(f)
        for r, v in c['lignes'].items():
            nom = valeur(c, v, 'nom')
            if not nom:
                continue
            existe, img = image_locale(valeur(c, v, 'thumbnail_url'))
            photo_ok, _ = image_locale(valeur(c, v, 'image_media'))
            lg, nb_liens = mesure_texte(valeur(c, v, 'texte_html'))
            morts, douteux = defauts_ligne(c, v)
            an = anomalies_ligne(c, v)
            out.append({
                'lg': lg, 'liens': nb_liens, 'morts': morts, 'douteux': douteux,
                'decede': any(a['type'] == 'decede' for a in an), 'ecart_dates': sum(1 for a in an if a['type'] != 'decede'),
                'f': f, 'l': r, 'nom': str(nom), 'slug': bd.slugify(str(nom)), 'dates': str(valeur(c, v, 'dates_affichage') or ''),
                'groupe': valeur(c, v, 'groupe') or '', 'image': img if existe else '',
                'image_prevue': img, 'video': bool(valeur(c, v, 'youtube_id')),
                'photo': bool(valeur(c, v, 'image_media')) and photo_ok,
                'annee': bd.to_int(valeur(c, v, 'annee_naissance')),
                **google_ligne(bd.slugify(str(nom))),
            })
    return out


def lire_fiche(f, l):
    c = classeur(f)
    if l not in c['lignes']:
        raise ValueError('Ligne introuvable')
    v = c['lignes'][l]
    champs = {h: ('' if v[i] is None else v[i]) for i, h in enumerate(c['entetes']) if h}
    existe, img = image_locale(champs.get('Thumbnail_URL'))
    photo_ok, photo = image_locale(champs.get('Image_media'))
    idx = index_sources()
    def source_de(url):
        m = cible_manquante(url)
        if not m or m[1].lower() not in idx:
            return None
        return {'source': os.path.relpath(idx[m[1].lower()], SOURCES_IMAGES).replace('\\', '/'),
                'cible': ('/pho/' if m[0] == PHO else '/photo/') + m[1] + '.webp'}
    return {'photo': photo, 'photo_existe': photo_ok,
            'source_portrait': source_de(champs.get('Thumbnail_URL')), 'source_photo': source_de(champs.get('Image_media')),'f': f, 'l': l, 'slug': bd.slugify(str(champs.get('Nom', ''))), 'entetes': [h for h in c['entetes'] if h], 'champs': champs,
            'image_existe': existe, 'image': img, 'nat_defaut': bd.NAT_DEFAUT.get(f, ''),
            'vivants': f in bd.FICHIERS_VIVANTS, 'courant': est_courant(f),
            'philosophes': rattaches().get(str(champs.get('Nom', '')).strip(), []) if est_courant(f) else [],
            'liens_defaut': liens.defauts_de(valeur(c, v, 'texte_html'), valeur(c, v, 'youtube_id')),
            'liens_controle': liens.etat().get('date', ''),
            'dates_anomalies': [] if est_courant(f) else anomalies_ligne(c, v),
            'dates_controle': dates.etat().get('date', ''),
            'google': {**statsgoogle.de_la_page(page_fiche(bd.slugify(str(champs.get('Nom', ''))), est_courant(f))),
                       'releve': statsgoogle.etat().get('date', ''), 'periode': statsgoogle.etat().get('periode', [])},
            'historique_propose': fichier_historique_propose(f, l) if f in bd.FICHIERS_VIVANTS else ''}


def controler_liens_fiche(f, l):
    c = classeur(f)
    v = c['lignes'][l]
    return liens.controler_fiche(valeur(c, v, 'texte_html'), valeur(c, v, 'youtube_id'))


# ── Courants : rattacher un philosophe, renommer partout ────────────────────
def rattacher(courant, f, l, retirer=False):
    """Ajoute (ou retire) le courant dans la colonne Courant d'un philosophe."""
    if f not in bd.PHILOSOPHE_FICHIERS:
        raise ValueError('Fichier de philosophe inconnu')
    courant = (courant or '').strip()
    c = classeur(f)
    if l not in c['lignes']:
        raise ValueError('Ligne introuvable')
    h = next((h for h in c['entetes'] if norm(h) == 'courant'), None)
    if not h:
        raise ValueError(f'{f} n\'a pas de colonne Courant')
    L = decouper_courants(valeur(c, c['lignes'][l], 'courant'))
    if retirer:
        L = [x for x in L if x != courant]
    elif courant not in L:
        L.append(courant)
    return ecrire(f, l, {h: ' ; '.join(L)})


def renommer_courant(ancien, nouveau):
    """Remplace l'ancien nom de courant par le nouveau dans la colonne Courant de tous les philosophes."""
    ancien, nouveau = ancien.strip(), nouveau.strip()
    n = 0
    for f in bd.PHILOSOPHE_FICHIERS:
        c = classeur(f)
        h = next((h for h in c['entetes'] if norm(h) == 'courant'), None)
        if not h:
            continue
        col = get_column_letter(c['entetes'].index(h) + 1)
        valeurs = {}
        for r, v in c['lignes'].items():
            L = decouper_courants(valeur(c, v, 'courant'))
            if ancien in L:
                valeurs[f'{col}{r}'] = ' ; '.join(nouveau if x == ancien else x for x in L)
        if valeurs:
            p = chemin(f)
            with verrou:
                if ouvert_dans_excel(p):
                    raise PermissionError(f'{f}.xlsx est ouvert dans Excel : fermez-le, le renommage n\'est pas fini.')
                sauvegarder(p)
                set_cells(p, valeurs)
                _cache.pop(f, None)
            n += len(valeurs)
    return n


# ── Écriture ───────────────────────────────────────────────────────────────
def ouvert_dans_excel(p):
    d, n = os.path.split(p)
    if os.path.exists(os.path.join(d, '~$' + n)):
        return True
    try:
        with open(p, 'r+b'):
            return False
    except PermissionError:
        return True


def sauvegarder(p):
    os.makedirs(SAUVEGARDES, exist_ok=True)
    horo = datetime.datetime.now().strftime('%Y%m%d-%H%M%S')
    shutil.copy2(p, os.path.join(SAUVEGARDES, f'{os.path.basename(p)}.{horo}'))


def typer(entete, v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    if s == '':
        return None
    if norm(entete) in ('annee_naissance', 'importance') and re.fullmatch(r'-?\d+', s):
        return int(s)
    return s


def ecrire(f, l, changements):
    p = chemin(f)
    with verrou:
        if ouvert_dans_excel(p):
            raise PermissionError(f'{f}.xlsx est ouvert dans Excel : fermez-le puis enregistrez de nouveau.')
        c = classeur(f)
        valeurs = {}
        for h, v in changements.items():
            if h not in c['entetes']:
                raise ValueError(f'Colonne inconnue : {h}')
            valeurs[f'{get_column_letter(c["entetes"].index(h) + 1)}{l}'] = typer(h, v)
        if not valeurs:
            return 0
        sauvegarder(p)
        set_cells(p, valeurs)
        _cache.pop(f, None)
        return len(valeurs)


# ── Ordre chronologique des lignes (année de naissance) ─────────────────────
LIGNE_RE = re.compile(r'<row [^>]*?r="(\d+)"[^>]*?(?:/>|>.*?</row>)', re.S)


def reordonner(p, ordre):
    """Réordonne les lignes de données (>= 4) de la 1re feuille sans réécrire le classeur.
    ordre : anciens numéros de ligne dans le nouvel ordre. Les cellules et les liens
    hypertexte (balises <hyperlink ref>) suivent leur ligne ; la mise en forme est conservée."""
    tmp = p + '.tmp'
    with zipfile.ZipFile(p) as zi:
        feuille = first_sheet_path(zi)
        xml = zi.read(feuille).decode('utf-8')
        lignes = [(int(m.group(1)), m) for m in LIGNE_RE.finditer(xml)]
        donnees = [(n, m) for n, m in lignes if n >= 4]
        numeros = sorted(n for n, _ in donnees)
        if sorted(ordre) != numeros:
            raise ValueError('Ordre incohérent')
        nouveau = dict(zip(ordre, numeros))  # ancien -> nouveau
        if all(a == b for a, b in nouveau.items()):
            return False
        par_num = {n: m.group(0) for n, m in donnees}

        def renum(bloc, a, b):
            bloc = re.sub(r'(<row [^>]*?\br=")%d(")' % a, r'\g<1>%d\2' % b, bloc, count=1)
            return re.sub(r'(<c r="[A-Z]+)%d(")' % a, r'\g<1>%d\2' % b, bloc)

        corps = ''.join(renum(par_num[a], a, nouveau[a]) for a in sorted(par_num, key=lambda a: nouveau[a]))
        debut, fin = donnees[0][1].start(), donnees[-1][1].end()
        xml = xml[:debut] + corps + xml[fin:]

        def lien(m):
            col, num = re.match(r'([A-Z]+)(\d+)$', m.group(2)).groups()
            return m.group(1) + col + str(nouveau.get(int(num), int(num))) + m.group(3)
        xml = re.sub(r'(<hyperlink [^>]*?\bref=")([A-Z]+\d+)(")', lien, xml)
        with zipfile.ZipFile(tmp, 'w') as zo:
            for it in zi.infolist():
                zo.writestr(it, xml.encode('utf-8') if it.filename == feuille else zi.read(it.filename))
    shutil.move(tmp, p)
    return True


def lignes_annees(f):
    """[(ligne, année ou None, a un nom)] dans l'ordre actuel du fichier."""
    c = classeur(f)
    p = chemin(f)
    with zipfile.ZipFile(p) as z:
        xml = z.read(first_sheet_path(z)).decode('utf-8')
    tous = [int(m.group(1)) for m in LIGNE_RE.finditer(xml) if int(m.group(1)) >= 4]
    out = []
    for n in tous:
        v = c['lignes'].get(n)
        out.append((n, bd.to_int(valeur(c, v, 'annee_naissance')) if v else None, bool(v and valeur(c, v, 'nom'))))
    return out


def ordre_trie(f):
    """Tri stable par année ; lignes sans nom (vides mises en forme) à la fin."""
    L = lignes_annees(f)
    pleines = [x for x in L if x[2]]
    vides = [x for x in L if not x[2]]
    sans_annee = 10 ** 9
    pleines_triees = sorted(pleines, key=lambda x: x[1] if x[1] is not None else sans_annee)
    return [x[0] for x in pleines_triees + vides]


def placer_ligne(f, l):
    """Déplace la seule ligne l à sa place chronologique : juste avant la première
    autre fiche née plus tard. Renvoie le nouveau numéro de la ligne."""
    L = lignes_annees(f)
    moi = next(x for x in L if x[0] == l)
    if moi[1] is None:
        return l
    autres = [x for x in L if x[0] != l]
    pos = next((i for i, x in enumerate(autres) if x[2] and x[1] is not None and x[1] > moi[1]), None)
    if pos is None:  # la plus récente : après la dernière fiche (avant les lignes vides)
        pos = max((i + 1 for i, x in enumerate(autres) if x[2]), default=0)
    ordre = [x[0] for x in autres]
    ordre.insert(pos, l)
    if ordre == [x[0] for x in L]:
        return l
    p = chemin(f)
    with verrou:
        if ouvert_dans_excel(p):
            raise PermissionError(f'{f}.xlsx est ouvert dans Excel : fermez-le.')
        reordonner(p, ordre)
        _cache.pop(f, None)
    numeros = sorted(x[0] for x in L)
    return numeros[ordre.index(l)]


def trier_fichier(f, simuler=False):
    L = lignes_annees(f)
    ordre = ordre_trie(f)
    bouge = sum(1 for a, b in zip([x[0] for x in L], ordre) if a != b)
    if simuler or not bouge:
        return bouge
    p = chemin(f)
    with verrou:
        if ouvert_dans_excel(p):
            raise PermissionError(f'{f}.xlsx est ouvert dans Excel : fermez-le.')
        sauvegarder(p)
        reordonner(p, ordre)
        _cache.pop(f, None)
    return bouge


# ── Nouvelle fiche ──────────────────────────────────────────────────────────
GROUPE_PAR_FICHIER = {
    'frise-philosophes-allemands': 'occident', 'frise-philosophes-americains': 'occident',
    'frise-philosophes-france': 'occident', 'frise-philosophes-france-actif': 'occident',
    'frise-philosophes-greco-romains': 'occident', 'frise-philosophes-medievaux': 'occident',
    'frise-philosophes-modernes': 'occident', 'frise-philosophes-renaissance-lumieres': 'occident',
    'frise-philosophes-russes': 'russe',
}


def dates_affichees(n, m):
    f = lambda y: f'{-y} A JC' if y < 0 else str(y)
    if m is None:
        return f(n)
    return f'{f(n)} - {f(m)}' if n < 0 else f'{f(n)} – {f(m)}'


def homonymes(nom, type_='philosophes'):
    cle = bd.slugify(nom)
    return [x for x in liste_fiches(type_) if x['slug'] == cle]


def creer_fiche(f, nom, annee, deces=None):
    nom = (nom or '').strip()
    if not nom:
        raise ValueError('Nom vide')
    n = bd.to_int(annee)
    if n is None:
        raise ValueError('Année de naissance invalide')
    m = bd.to_int(deces) if str(deces or '').strip() else None
    p = chemin(f)
    c = classeur(f)
    L = lignes_annees(f)
    pleines = [x[0] for x in L if x[2]]
    vides = [x[0] for x in L if not x[2]]
    ligne = vides[0] if vides else (max([x[0] for x in L] or [3]) + 1)
    modele = pleines[-1] if pleines else 4
    lien = 'https://fr.wikipedia.org/wiki/' + urllib.parse.quote(nom.replace(' ', '_'))
    valeurs_par_cle = {
        'annee_naissance': n, 'annee_deces': None if m is None else str(m),
        'dates_affichage': dates_affichees(n, m), 'nom': nom,
        'texte_html': f'<p><a href="{lien}" title=" Wikipedia">{nom}</a> </p>',
        'nationalite': bd.NAT_DEFAUT.get(f) or None,
        'groupe': GROUPE_PAR_FICHIER.get(f),
    }
    if est_courant(f):  # un courant : une seule année (« 360 av. J.-C. »), groupe Occident pour le fichier occidental
        valeurs_par_cle['groupe'] = 'occident' if f == 'frise-courant-pensee-occidental' else None
        valeurs_par_cle['annee_deces'] = None
        valeurs_par_cle['dates_affichage'] = f'{-n} av. J.-C.' if n < 0 else str(n)
    valeurs, styles = {}, {}
    for i, h in enumerate(c['entetes']):
        col = get_column_letter(i + 1)
        ref = f'{col}{ligne}'
        styles[ref] = f'{col}{modele}'
        valeurs[ref] = valeurs_par_cle.get(norm(h)) if norm(h) in valeurs_par_cle else None
    with verrou:
        if ouvert_dans_excel(p):
            raise PermissionError(f'{f}.xlsx est ouvert dans Excel : fermez-le puis recommencez.')
        sauvegarder(p)
        set_cells(p, valeurs, style_from=styles)
        _cache.pop(f, None)
    return placer_ligne(f, ligne)


def enregistrer_image(f, l, nom, data_url, cadrage, tel_quel=False):
    nom = re.sub(r'[^a-z0-9-]', '', nom.lower().replace(' ', '-')).strip('-')
    if not nom:
        raise ValueError("Nom d'image vide")
    brut = base64.b64decode(data_url.split(',', 1)[1])
    im = ImageOps.exif_transpose(Image.open(io.BytesIO(brut))).convert('RGB')
    # Portrait 5:6 ; cadrage = position verticale du recadrage (0 = haut, 1 = bas).
    # Image déjà prête (petite, recadrée à la main) : convertie telle quelle.
    w, h = im.size
    if tel_quel and max(w, h) <= 160:
        return _sauver_mini(f, l, nom, im)
    cible = 5 / 6
    if w / h > cible:
        nw = int(h * cible)
        x0 = (w - nw) // 2
        im = im.crop((x0, 0, x0 + nw, h))
    else:
        nh = int(w / cible)
        y0 = int((h - nh) * max(0.0, min(1.0, float(cadrage))))
        im = im.crop((0, y0, w, y0 + nh))
    im = im.resize((100, 120), Image.LANCZOS)
    return _sauver_mini(f, l, nom, im)


def _sauver_mini(f, l, nom, im):
    fichier = f'{nom}-mini.webp'
    dest = os.path.join(PHO, fichier)
    if os.path.exists(dest):
        os.makedirs(SAUVEGARDES, exist_ok=True)
        shutil.copy2(dest, os.path.join(SAUVEGARDES, fichier + datetime.datetime.now().strftime('.%Y%m%d-%H%M%S')))
    im.save(dest, 'WEBP', quality=85, method=6)
    ecrire(f, l, {'Thumbnail_URL': PREFIXE_IMAGE + fichier})
    return '/pho/' + fichier


def enregistrer_photo(f, l, nom, data_url):
    """Grande photo (à la place d'une vidéo) : non recadrée, réduite à 400 px de haut au plus."""
    nom = re.sub(r'[^a-z0-9-]', '', nom.lower().replace(' ', '-')).strip('-')
    if not nom:
        raise ValueError('Nom de photo vide')
    brut = base64.b64decode(data_url.split(',', 1)[1])
    im = ImageOps.exif_transpose(Image.open(io.BytesIO(brut))).convert('RGB')
    im.thumbnail((700, 400), Image.LANCZOS)
    fichier = f'{nom}.webp'
    dest = os.path.join(PHOTO, fichier)
    if os.path.exists(dest):
        os.makedirs(SAUVEGARDES, exist_ok=True)
        shutil.copy2(dest, os.path.join(SAUVEGARDES, fichier + datetime.datetime.now().strftime('.%Y%m%d-%H%M%S')))
    im.save(dest, 'WEBP', quality=85, method=6)
    ecrire(f, l, {'Image_media': PREFIXE_PHOTO + fichier})
    return '/photo/' + fichier


# ── Images en attente : jpg de pho-jpg dont une fiche attend la version webp ─
def index_sources():
    """nom (sans extension, minuscules) -> chemin du fichier source, dans pho-jpg et ses sous-dossiers."""
    idx = {}
    if not os.path.isdir(SOURCES_IMAGES):
        return idx
    for racine, dossiers, fichiers in os.walk(SOURCES_IMAGES):
        dossiers[:] = [d for d in dossiers if not d.startswith('sca-')]
        for f in fichiers:
            base, ext = os.path.splitext(f)
            if ext.lower() not in ('.jpg', '.jpeg', '.png', '.webp'):
                continue
            base = re.sub(r'\.webp$', '', base, flags=re.I).lower()
            p = os.path.join(racine, f)
            if base not in idx or os.path.getmtime(p) > os.path.getmtime(idx[base]):
                idx[base] = p
    return idx


def cible_manquante(url):
    """('/pho/x-mini.webp' ou '/photo/x.webp') -> (dossier, nom de base) si le webp manque, sinon None."""
    s = bd.url_relative(url or '')
    for prefixe, dossier in (('/pho/', PHO), ('/photo/', PHOTO)):
        if s.startswith(prefixe) and s.lower().endswith('.webp'):
            nom = s[len(prefixe):].lstrip('/')
            if not os.path.exists(os.path.join(dossier, nom)):
                return dossier, nom[:-5]
    return None


def images_en_attente():
    idx = index_sources()
    trouvees, sans_source = {}, []
    for f in bd.PHILOSOPHE_FICHIERS:
        c = classeur(f)
        for r, v in c['lignes'].items():
            nom = valeur(c, v, 'nom')
            for col, genre in (('thumbnail_url', 'portrait'), ('image_media', 'grande photo')):
                m = cible_manquante(valeur(c, v, col))
                if not m:
                    continue
                dossier, base = m
                src = idx.get(base.lower())
                if src:
                    cle = os.path.join(dossier, base + '.webp')
                    e = trouvees.setdefault(cle, {'source': os.path.relpath(src, SOURCES_IMAGES).replace('\\', '/'),
                                                 'cible': ('/pho/' if dossier == PHO else '/photo/') + base + '.webp',
                                                 'genre': genre, 'fiches': []})
                    e['fiches'].append(str(nom))
                else:
                    sans_source.append({'nom': str(nom), 'attendu': base + '.webp', 'genre': genre})
    return {'images': sorted(trouvees.values(), key=lambda x: x['cible']), 'sans_source': sans_source,
            'dossier': SOURCES_IMAGES}


def convertir_source(source, cible):
    """Convertit un jpg (déjà recadré) en webp, sans le recadrer ; une grande photo est limitée à 700×400."""
    src = os.path.normpath(os.path.join(SOURCES_IMAGES, source))
    if not src.startswith(os.path.normpath(SOURCES_IMAGES)) or not os.path.exists(src):
        raise ValueError(f'Source introuvable : {source}')
    m = re.fullmatch(r'/(pho|photo)/([a-z0-9._-]+\.webp)', cible, re.I)
    if not m:
        raise ValueError(f'Cible invalide : {cible}')
    im = ImageOps.exif_transpose(Image.open(src)).convert('RGB')
    if m.group(1) == 'photo':
        im.thumbnail((700, 400), Image.LANCZOS)
    elif max(im.size) > 160:  # portrait trop grand pour une vignette : on le réduit (sans recadrer)
        im.thumbnail((100, 120), Image.LANCZOS)
    dest = os.path.join(PHO if m.group(1) == 'pho' else PHOTO, m.group(2))
    im.save(dest, 'WEBP', quality=82, method=6)
    return cible


# ── Sites qui refusent l'affichage dans le panneau (ouverts en petite fenêtre) ─
def lire_domaines():
    return json.load(open(DOMAINES, encoding='utf-8'))


def ecrire_domaines(liste):
    """Les entrées déjà présentes sont gardées telles quelles ; seules les nouvelles sont nettoyées."""
    actuelles = lire_domaines()
    propre = []
    for d in liste:
        if d in actuelles:
            propre.append(d)
            continue
        d = re.sub(r'^https?://', '', str(d).strip().lower()).split('/')[0]
        d = re.sub(r'^www\.', '', d)
        if d and re.fullmatch(r'[a-z0-9.-]+\.[a-z]{2,}', d) and d not in propre:
            propre.append(d)
    os.makedirs(SAUVEGARDES, exist_ok=True)
    shutil.copy2(DOMAINES, os.path.join(SAUVEGARDES, 'iframe-blocked-domains.json' + datetime.datetime.now().strftime('.%Y%m%d-%H%M%S')))
    with open(DOMAINES, 'w', encoding='utf-8', newline='\n') as f:
        f.write(json.dumps(propre, ensure_ascii=False, indent=2) + '\n')
    return propre


# ── YouTube ────────────────────────────────────────────────────────────────
def id_youtube(s):
    s = (s or '').strip()
    if re.fullmatch(r'[\w-]{11}', s):
        return s
    m = re.search(r'(?:v=|youtu\.be/|embed/|shorts/|live/)([\w-]{11})', s)
    return m.group(1) if m else ''


def infos_youtube(s):
    vid = id_youtube(s)
    if not vid:
        return {'ok': False, 'message': "Lien YouTube non reconnu"}
    url = 'https://www.youtube.com/oembed?format=json&url=' + urllib.parse.quote(f'https://www.youtube.com/watch?v={vid}')
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=10) as r:
            d = json.load(r)
        titre, chaine = d.get('title', ''), d.get('author_name', '')
        return {'ok': True, 'id': vid, 'titre': titre, 'chaine': chaine,
                'credit': f'"{titre}" par {chaine}' if chaine else titre}
    except urllib.error.HTTPError as e:
        msg = ('Vidéo introuvable ou supprimée' if e.code in (400, 404) else
               "Vidéo privée, ou intégration désactivée par son auteur : elle ne s'affichera pas sur le site" if e.code in (401, 403) else
               f'Erreur YouTube {e.code}')
        return {'ok': False, 'id': vid, 'message': msg}
    except Exception as e:  # pas de connexion…
        return {'ok': True, 'id': vid, 'titre': '', 'chaine': '', 'credit': '', 'message': f'Vérification impossible ({e.__class__.__name__})'}


# ── Liens du texte : vérification d'une adresse (Wikipédia, YouTube, autre site) ──
def infos_lien(url):
    """{'ok': True/False/None (doute), 'genre', 'titre' (titre proposé pour l'attribut title), 'message'}"""
    url = (url or '').strip()
    if id_youtube(url) and re.search(r'youtu\.?be', url):
        y = infos_youtube(url)
        titre = f"{y.get('titre', '')} par {y['chaine']}" if y.get('chaine') else y.get('titre', '')
        return {'ok': y['ok'], 'genre': 'youtube', 'titre': titre, 'message': y.get('message', '')}
    m = re.match(r'https?://([a-z-]+)\.(?:m\.)?wikipedia\.org/wiki/([^#?]+)', url)
    if m:
        lang, page = m.groups()
        api = f'https://{lang}.wikipedia.org/api/rest_v1/page/summary/' + urllib.parse.quote(urllib.parse.unquote(page), safe='')
        try:
            req = urllib.request.Request(api, headers={'User-Agent': 'laphilo-editeur/1.0 (contact@laphilo.fr)'})
            with urllib.request.urlopen(req, timeout=10) as r:
                d = json.load(r)
            titre = ' Wikipedia' if lang == 'fr' else f' Wikipedia ({lang})'
            if d.get('type') == 'disambiguation':
                return {'ok': None, 'genre': 'wikipedia', 'titre': titre, 'message': "page d'homonymie : choisir la bonne page"}
            canon = d.get('titles', {}).get('normalized', '')
            return {'ok': True, 'genre': 'wikipedia', 'titre': titre, 'message': f'« {canon} »' if canon else ''}
        except urllib.error.HTTPError as e:
            return {'ok': False if e.code == 404 else None, 'genre': 'wikipedia', 'titre': '',
                    'message': 'page introuvable' if e.code == 404 else f'Wikipédia répond {e.code}'}
        except Exception as e:
            return {'ok': None, 'genre': 'wikipedia', 'titre': '', 'message': f'vérification impossible ({e.__class__.__name__})'}
    if not url.startswith('http'):
        return {'ok': False, 'genre': 'autre', 'titre': '', 'message': "l'adresse ne commence pas par http"}
    for methode in ('HEAD', 'GET'):
        try:
            req = urllib.request.Request(url, method=methode, headers={'User-Agent': 'Mozilla/5.0'})
            with urllib.request.urlopen(req, timeout=10) as r:
                return {'ok': True, 'genre': 'autre', 'titre': '', 'message': ''}
        except urllib.error.HTTPError as e:
            if methode == 'HEAD' and e.code in (403, 405, 429, 501):
                continue  # certains sites refusent HEAD : on réessaie en GET
            return {'ok': False if e.code in (404, 410) else None, 'genre': 'autre', 'titre': '',
                    'message': f'page introuvable ({e.code})' if e.code in (404, 410) else f'le site répond {e.code} (il refuse peut-être les vérifications automatiques)'}
        except Exception as e:
            if methode == 'HEAD':
                continue
            return {'ok': None, 'genre': 'autre', 'titre': '', 'message': f'vérification impossible ({e.__class__.__name__})'}


# ── Audio France Culture (lien de la page de l'émission ou du fichier) ──────
def infos_audio(url):
    import html as html_mod
    url = (url or '').strip()
    if not url.startswith('http'):
        return {'ok': False, 'message': 'Collez un lien commençant par https://'}
    if re.search(r'\.(mp3|m4a)(\?|$)', url, re.I):
        return {'ok': True, 'src': url, 'titre': ''}
    try:
        req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=15) as r:
            page = r.read().decode('utf-8', errors='replace')
    except Exception as e:
        return {'ok': False, 'message': f'Page inaccessible ({e.__class__.__name__})'}
    m = (re.search(r'"contentUrl"\s*:\s*"(https://media\.radiofrance-podcast\.net/[^"]+?\.(?:mp3|m4a))"', page)
         or re.search(r'(https://media\.radiofrance-podcast\.net/[^"\s\]+?\.(?:mp3|m4a))', page))
    if not m:
        return {'ok': False, 'message': "Aucun fichier audio trouvé sur cette page (émission pas encore en podcast ?)"}
    t = re.search(r'<meta property="og:title" content="([^"]*)"', page)
    titre = html_mod.unescape(t.group(1)) if t else ''
    titre = re.split(r'\s+:\s+épisode|\s+\|\s+', titre)[0].strip()
    radio = 'France Culture' if 'franceculture' in url else 'Radio France'
    return {'ok': True, 'src': m.group(1), 'titre': f'{titre} par {radio}' if titre else ''}


# ── Codes valides (menus) ─────────────────────────────────────────────────
def codes():
    P = json.load(open(os.path.join(BASE, 'src', 'data', 'philosophes.json'), encoding='utf-8'))
    C = json.load(open(os.path.join(BASE, 'src', 'data', 'courants.json'), encoding='utf-8'))
    branches_c = []
    for c in C:
        for b in c.get('branches', []):
            if b not in branches_c:
                branches_c.append(b)
    return {
        'fichiers': [{'f': f, 'label': lab} for f, (_, lab) in bd.PHILOSOPHE_FICHIERS.items()],
        'fichiers_courants': [{'f': f, 'label': lab} for f, (_, lab) in bd.COURANT_FICHIERS.items()],
        'branches_courants': sorted(branches_c, key=lambda b: bd.normalize_header(b)),
        'groupes': [{'code': k, 'label': v} for k, v in bd.GROUPES.items()],
        'branches': BRANCHES, 'traditions': TRADITIONS,
        'themes': [{'code': k, 'label': v} for k, v in bd.THEMES.items()],
        'courants': sorted({c['name'] for c in C}, key=str.lower),
        # familles déjà utilisées, par nature (mouvement : familles historiques ; position : grandes questions)
        'familles': {n: sorted({c['famille'] for c in C if c.get('nature') == n and c.get('famille')}, key=str.lower)
                     for n in ('mouvement', 'position')},
        'photos': sorted(os.path.splitext(n)[0] for n in os.listdir(PHOTO) if n.endswith('.webp')),
        'nationalites': sorted({p.get('nationalite') for p in P if p.get('nationalite')}),
    }


def lancer_build():
    r = subprocess.run([sys.executable, os.path.join(BASE, 'scripts', 'build-data.py')], cwd=BASE,
                       capture_output=True, text=True, encoding='utf-8', errors='replace',
                       env={**os.environ, 'PYTHONIOENCODING': 'utf-8'})
    sortie = (r.stdout or '') + (r.stderr or '')
    i = sortie.find('=== Vérifications')
    return {'ok': r.returncode == 0, 'resume': sortie[i:] if i >= 0 else sortie[-3000:]}


# ── Aperçu local : le site Astro (astro dev) à jour des xlsx ────────────────
SITE_LOCAL = 'http://localhost:4321'
verrou_apercu = threading.Lock()


def site_local_actif():
    try:
        with urllib.request.urlopen(SITE_LOCAL + '/', timeout=3) as r:
            return r.status == 200
    except Exception:
        return False


def astro(*args):
    subprocess.run('npx astro dev ' + ' '.join(args), cwd=BASE, shell=True,
                   capture_output=True, text=True, encoding='utf-8', errors='replace', timeout=120)


def preparer_apercu(slug, type_='philosophes'):
    """Régénère les données si un xlsx a changé, (re)lance le site local, attend que la fiche réponde."""
    with verrou_apercu:
        json_phi = os.path.join(BASE, 'src', 'data', 'philosophes.json')
        dernier_xlsx = max(os.path.getmtime(p) for p in glob.glob(os.path.join(DATA, 'frise-*.xlsx')))
        photos = [os.path.getmtime(p) for d in (PHO, PHOTO) for p in glob.glob(os.path.join(d, '*.webp'))]
        perime = max([dernier_xlsx] + photos) > os.path.getmtime(json_phi)
        etapes = []
        if perime:
            r = lancer_build()
            etapes.append('données régénérées')
            if not r['ok']:
                return {'ok': False, 'message': 'build-data a échoué', 'detail': r['resume']}
        actif = site_local_actif()
        if actif and perime:
            astro('stop')  # le site local garde sinon les anciennes données en mémoire
            actif = False
            etapes.append('site local redémarré')
        if not actif:
            astro('--background')
            if 'site local redémarré' not in etapes:
                etapes.append('site local lancé')
        url = f'{SITE_LOCAL}/{"courants" if type_ == "courants" else "philosophes"}/{slug}/'
        import time
        for _ in range(90):
            try:
                with urllib.request.urlopen(url, timeout=20) as r:
                    if r.status == 200:
                        return {'ok': True, 'url': url, 'message': ', '.join(etapes) or 'déjà à jour'}
            except urllib.error.HTTPError as e:
                if e.code == 404:
                    return {'ok': False, 'url': url, 'message': "Page introuvable sur le site local (nom modifié ? fiche nouvelle ?)"}
            except Exception:
                pass
            time.sleep(1)
        return {'ok': False, 'url': url, 'message': "Le site local ne répond pas (lancez « npx astro dev » à la main)"}


PAGE_APERCU = '''<!doctype html><html lang="fr"><head><meta charset="utf-8"><title>Aperçu local…</title>
<style>body{margin:0;height:100vh;display:grid;place-items:center;background:#1a0f07;color:#f6ecd0;font-family:Georgia,serif;text-align:center}
.r{width:46px;height:46px;border:4px solid #5c4318;border-top-color:#d4a843;border-radius:50%;margin:0 auto 1.2rem;animation:t 1s linear infinite}
@keyframes t{to{transform:rotate(360deg)}} small{color:#ead9a8;display:block;margin-top:.6rem;max-width:520px} pre{text-align:left;white-space:pre-wrap;font-size:.8rem;color:#ffb4a8}</style></head>
<body><div><div class="r" id="r"></div><div id="m">Préparation de l'aperçu local…</div>
<small>Si des fiches ont été modifiées, les données sont régénérées et le site local redémarré (10 à 30 secondes).</small><pre id="d"></pre></div>
<script>fetch('/api/preparer?slug=__SLUG__&type=__TYPE__').then(r=>r.json()).then(d=>{if(d.ok){location.replace(d.url)}else{
document.getElementById('r').remove();document.getElementById('m').textContent='⚠ '+d.message;document.getElementById('d').textContent=d.detail||''}})
.catch(e=>{document.getElementById('m').textContent='⚠ '+e})</script></body></html>'''


# ── Serveur HTTP ───────────────────────────────────────────────────────────
class Gestion(BaseHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def envoyer(self, code, corps, type_='application/json; charset=utf-8'):
        if not isinstance(corps, bytes):
            corps = json.dumps(corps, ensure_ascii=False).encode('utf-8')
        self.send_response(code)
        self.send_header('Content-Type', type_)
        self.send_header('Cache-Control', 'no-store')
        self.end_headers()
        self.wfile.write(corps)

    def do_GET(self):
        u = urllib.parse.urlparse(self.path)
        q = dict(urllib.parse.parse_qsl(u.query))
        try:
            if u.path in ('/', '/index.html'):
                if relancer_si_mis_a_jour():
                    return self.envoyer(200, PAGE_RELANCE.encode('utf-8'), 'text/html; charset=utf-8')
                return self.envoyer(200, open(os.path.join(ICI, 'index.html'), 'rb').read(), 'text/html; charset=utf-8')
            if u.path == '/favicon.ico':  # icône de la fenêtre et de la barre des tâches
                return self.envoyer(200, open(os.path.join(ICI, 'editeur.ico'), 'rb').read(), 'image/x-icon')
            if u.path.startswith('/source/'):
                p = os.path.normpath(os.path.join(SOURCES_IMAGES, urllib.parse.unquote(u.path[len('/source/'):])))
                if p.startswith(os.path.normpath(SOURCES_IMAGES)) and os.path.exists(p):
                    return self.envoyer(200, open(p, 'rb').read(), 'image/png' if p.lower().endswith('.png') else 'image/jpeg')
                return self.envoyer(404, {'erreur': 'image absente'})
            if u.path == '/api/images-attente':
                return self.envoyer(200, images_en_attente())
            if u.path == '/api/domaines':
                return self.envoyer(200, lire_domaines())
            if u.path.startswith('/pho/') or u.path.startswith('/photo/'):
                p = os.path.join(PHO if u.path.startswith('/pho/') else PHOTO, os.path.basename(urllib.parse.unquote(u.path)))
                if os.path.exists(p):
                    return self.envoyer(200, open(p, 'rb').read(), 'image/webp' if p.endswith('.webp') else 'image/jpeg')
                return self.envoyer(404, {'erreur': 'image absente'})
            if u.path.startswith('/apercu/'):  # /apercu/<slug>/ (philosophe) ou /apercu/courants/<slug>/
                reste = u.path[len('/apercu/'):].strip('/')
                type_ = 'courants' if reste.startswith('courants/') else 'philosophes'
                slug = re.sub(r'[^a-z0-9-]', '', reste.split('/')[-1])
                return self.envoyer(200, PAGE_APERCU.replace('__SLUG__', slug).replace('__TYPE__', type_).encode('utf-8'), 'text/html; charset=utf-8')
            if u.path == '/api/preparer':
                return self.envoyer(200, preparer_apercu(re.sub(r'[^a-z0-9-]', '', q.get('slug', '')), q.get('type', 'philosophes')))
            if u.path == '/api/version':
                return self.envoyer(200, {'version': VERSION})
            if u.path == '/api/codes':
                return self.envoyer(200, codes())
            if u.path == '/api/fiches':
                return self.envoyer(200, liste_fiches(q.get('type', 'philosophes')))
            if u.path == '/api/fiche':
                return self.envoyer(200, lire_fiche(q['f'], int(q['l'])))
            if u.path == '/api/audio':
                return self.envoyer(200, infos_audio(q.get('url', '')))
            if u.path == '/api/lien':
                return self.envoyer(200, infos_lien(q.get('url', '')))
            if u.path == '/api/youtube':
                return self.envoyer(200, infos_youtube(q.get('url', '')))
            if u.path == '/api/dates':  # écarts avec Wikidata + état du contrôle
                return self.envoyer(200, liste_dates() if q.get('fiches') else dates.resume())
            if u.path == '/api/google':  # chiffres Search Console et indexation des fiches
                return self.envoyer(200, liste_google() if q.get('fiches') else statsgoogle.resume())
            if u.path == '/api/publication':  # état de la publication en cours (suivi des étapes)
                return self.envoyer(200, publier.suivi_public())
            if u.path == '/api/liens-morts':  # fiches en défaut + état du contrôle complet
                return self.envoyer(200, liste_liens_morts() if q.get('fiches') else liens.resume())
            return self.envoyer(404, {'erreur': 'introuvable'})
        except Exception as e:
            return self.envoyer(400, {'erreur': str(e)})

    def do_POST(self):
        u = urllib.parse.urlparse(self.path)
        try:
            n = int(self.headers.get('Content-Length', 0))
            d = json.loads(self.rfile.read(n) or b'{}')
            if u.path == '/api/fiche':
                renommes = 0
                if est_courant(d['f']):  # nom de courant changé : les philosophes qui le citent suivent
                    c = classeur(d['f'])
                    h_nom = next((h for h in c['entetes'] if norm(h) == 'nom'), None)
                    ancien = str(valeur(c, c['lignes'].get(int(d['l']), []), 'nom') or '').strip() if int(d['l']) in c['lignes'] else ''
                    nouveau = str(d['changements'].get(h_nom, '')).strip() if h_nom in d['changements'] else ''
                nb = ecrire(d['f'], int(d['l']), d['changements'])
                if est_courant(d['f']) and ancien and nouveau and nouveau != ancien:
                    renommes = renommer_courant(ancien, nouveau)
                l = placer_ligne(d['f'], int(d['l']))
                return self.envoyer(200, {'ok': True, 'cellules': nb, 'ligne': l, 'deplacee': l != int(d['l']), 'renommes': renommes})
            if u.path == '/api/rattacher':
                rattacher(d['courant'], d['f'], int(d['l']), d.get('retirer', False))
                return self.envoyer(200, {'ok': True, 'philosophes': rattaches().get(d['courant'].strip(), [])})
            if u.path == '/api/nouvelle':
                if not d.get('confirme'):
                    h = homonymes(d.get('nom', ''), 'courants' if est_courant(d['f']) else 'philosophes')
                    if h:
                        return self.envoyer(200, {'ok': False, 'homonymes': h})
                l = creer_fiche(d['f'], d.get('nom'), d.get('annee'), d.get('deces'))
                return self.envoyer(200, {'ok': True, 'f': d['f'], 'ligne': l})
            if u.path == '/api/trier':
                n = trier_fichier(d['f'], simuler=d.get('simuler', False))
                return self.envoyer(200, {'ok': True, 'lignes': n})
            if u.path == '/api/image':
                url = enregistrer_image(d['f'], int(d['l']), d['nom'], d['data'], d.get('cadrage', 0.15), d.get('tel_quel', False))
                return self.envoyer(200, {'ok': True, 'image': url})
            if u.path == '/api/images-importer':
                faites, erreurs = [], []
                for x in d.get('images', []):
                    try:
                        faites.append(convertir_source(x['source'], x['cible']))
                    except Exception as e:
                        erreurs.append(f"{x.get('source')} : {e}")
                return self.envoyer(200, {'ok': not erreurs, 'faites': faites, 'erreurs': erreurs})
            if u.path == '/api/domaines':
                return self.envoyer(200, {'ok': True, 'domaines': ecrire_domaines(d.get('domaines', []))})
            if u.path == '/api/photo':
                url = enregistrer_photo(d['f'], int(d['l']), d['nom'], d['data'])
                return self.envoyer(200, {'ok': True, 'photo': url})
            if u.path == '/api/arret':  # demandé par un éditeur plus récent qui démarre
                self.envoyer(200, {'ok': True})
                threading.Thread(target=self.server.shutdown, daemon=True).start()
                return
            if u.path == '/api/dates-controle':
                return self.envoyer(200, {'ok': dates.lancer_controle(philosophes_slug_texte), **dates.resume()})
            if u.path == '/api/dates-ignorer':  # la fiche a raison : écart plus signalé
                dates.ignorer(d['cle'])
                return self.envoyer(200, {'ok': True})
            if u.path == '/api/google-controle':  # chiffres + indexation (quota du jour), en arrière-plan
                ok = statsgoogle.lancer_controle(toutes_les_pages) if d.get('index') else statsgoogle.lancer_chiffres()
                return self.envoyer(200, {'ok': ok, **statsgoogle.resume()})
            if u.path == '/api/deplacer':  # philosophe décédé : fichier « actif » -> fichier historique
                l, perdues = deplacer_fiche(d['f'], int(d['l']), d['cible'], d.get('deces'))
                return self.envoyer(200, {'ok': True, 'f': d['cible'], 'ligne': l, 'perdues': perdues})
            if u.path == '/api/publication-preparer':  # build-data + ce qui partira
                return self.envoyer(200, publier.preparer())
            if u.path == '/api/publier':
                return self.envoyer(200, {'ok': publier.publier(d.get('message', ''), d.get('relance', False))})
            if u.path == '/api/liens-controle':  # contrôle complet, en arrière-plan
                return self.envoyer(200, {'ok': liens.lancer_controle(toutes_les_fiches), **liens.resume()})
            if u.path == '/api/liens-fiche':  # après un enregistrement : adresses nouvelles ou en défaut
                return self.envoyer(200, {'ok': True, 'liens_defaut': controler_liens_fiche(d['f'], int(d['l']))})
            if u.path == '/api/liens-bon':  # « à vérifier » ouvert à la main et trouvé bon
                liens.marquer_bon(d['url'])
                return self.envoyer(200, {'ok': True})
            if u.path == '/api/build':
                return self.envoyer(200, lancer_build())
            return self.envoyer(404, {'erreur': 'introuvable'})
        except PermissionError as e:
            return self.envoyer(409, {'erreur': str(e)})
        except Exception as e:
            return self.envoyer(400, {'erreur': f'{e.__class__.__name__} : {e}'})


def ouvrir_fenetre(adresse):
    """Ouvre l'éditeur comme une application : fenêtre Edge (ou Chrome) en mode « app », sans barre d'adresse
    ni onglets, avec son icône dans la barre des tâches. Son profil à part retient la taille et la position.
    Option --navigateur : onglet du navigateur habituel (comme avant). L'adresse reste utilisable dans tout navigateur."""
    if '--navigateur' not in sys.argv:
        pf = [os.environ.get(v, '') for v in ('ProgramFiles(x86)', 'ProgramFiles', 'LOCALAPPDATA')]
        candidats = [os.path.join(b, r) for b in pf if b for r in (
            r'Microsoft\Edge\Application\msedge.exe', r'Google\Chrome\Application\chrome.exe')]
        exe = next((c for c in candidats if os.path.exists(c)), None)
        if exe:
            profil = os.path.join(os.environ.get('LOCALAPPDATA', ICI), 'LaPhilo', 'editeur-fenetre')
            try:
                subprocess.Popen([exe, f'--app={adresse}', f'--user-data-dir={profil}', '--window-size=1500,950',
                                  '--no-first-run', '--no-default-browser-check'])
                return
            except OSError:
                pass
    webbrowser.open(adresse)


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    if '--port' in sys.argv:  # pour un essai à côté d'un éditeur déjà ouvert
        PORT = int(sys.argv[sys.argv.index('--port') + 1])
    adresse = f'http://127.0.0.1:{PORT}/'
    class Serveur(ThreadingHTTPServer):
        # Sous Windows, SO_REUSEADDR laisse un second éditeur ouvrir le même port sans erreur :
        # l'ancien n'était alors jamais remplacé et continuait à répondre avec son ancien code.
        allow_reuse_address = False
        allow_reuse_port = False

        def server_bind(self):
            import socket
            if hasattr(socket, 'SO_EXCLUSIVEADDRUSE'):
                self.socket.setsockopt(socket.SOL_SOCKET, socket.SO_EXCLUSIVEADDRUSE, 1)
            super().server_bind()

    def demarrer():
        return Serveur(('127.0.0.1', PORT), Gestion)
    try:
        serveur = demarrer()
    except OSError:
        # Déjà lancé : même version -> on rouvre la page ; version plus ancienne -> on la remplace
        try:
            with urllib.request.urlopen(adresse + 'api/version', timeout=3) as r:
                meme = json.load(r).get('version') == VERSION
        except Exception:
            meme = False
        if meme:
            ouvrir_fenetre(adresse)
            sys.exit(0)
        try:
            urllib.request.urlopen(urllib.request.Request(adresse + 'api/arret', data=b'{}', method='POST'), timeout=3)
        except Exception:
            pass
        import time
        for _ in range(20):
            time.sleep(0.3)
            try:
                serveur = demarrer()
                break
            except OSError:
                serveur = None
        if serveur is None:
            print("Un ancien éditeur est encore ouvert : fermez sa fenêtre noire, puis relancez.")
            ouvrir_fenetre(adresse)
            sys.exit(1)
    print(f'Éditeur des fiches : {adresse}  (fermez cette fenêtre pour arrêter)')
    if '--sans-navigateur' not in sys.argv:
        threading.Timer(0.8, lambda: ouvrir_fenetre(adresse)).start()
    try:
        serveur.serve_forever()
    except KeyboardInterrupt:
        pass
